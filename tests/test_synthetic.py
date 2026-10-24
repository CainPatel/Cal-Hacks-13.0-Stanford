"""Rendered documents carry exactly the values in truth.json, in the style it names."""
import json
from decimal import Decimal
from email import message_from_bytes

import pdfplumber
import pytest
from openpyxl import load_workbook

from synthetic import render
from synthetic.__main__ import main


@pytest.fixture(scope="module")
def out(tmp_path_factory):
    out = tmp_path_factory.mktemp("synthetic")
    assert main(out) == 15 + 25
    return out


def folders(out):
    return sorted(p for p in out.glob("*/*") if p.is_dir())


def test_pdf_and_xlsx_match_truth(out):
    styles = set()
    for f in folders(out):
        truth = json.loads((f / "truth.json").read_text())
        style, req = truth["style"], truth["pay_request"]
        styles.add(style)
        with pdfplumber.open(f / "g703.pdf") as pdf:
            text = pdf.pages[0].extract_text()
        ws = load_workbook(f / "g703.xlsx").active
        rows = {r[0]: r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]}
        for line in req["lines"]:
            for field in ("scheduled_value", "total_to_date", "retainage"):
                assert render.money(Decimal(line[field]), style) in text, (f.name, field)
            assert render.pct(Decimal(line["pct_complete"]), style) in text, f.name
            assert Decimal(str(rows[line["item_no"]][6])) == Decimal(line["total_to_date"])
        assert render.money(Decimal(req["g702"]["current_payment_due"]), style) in text
        assert req["bank_account"] in text
    assert styles == set(render.STYLES)  # every style shows up somewhere


def test_documents_follow_the_anomaly(out):
    cases = {f.name: f for f in (out / "cases").iterdir()}
    missing = [f for n, f in cases.items() if "missing_conditional_release" in n]
    assert sum(not (f / "release.pdf").exists() for f in missing) == 2  # the third attaches a stale one
    homoglyph = next(f for n, f in cases.items() if "homoglyph" in n)
    [eml] = homoglyph.glob("email_*.eml")
    sender = message_from_bytes(eml.read_bytes())["From"]
    vendor = json.loads((homoglyph / "truth.json").read_text())["vendor"]
    assert sender.split("@")[1] != vendor["email_domain"]
    assert all((f / "coi.pdf").exists() and (f / "g703.pdf").exists() for f in folders(out))
